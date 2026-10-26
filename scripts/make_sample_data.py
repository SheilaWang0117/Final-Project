"""Generate import templates and SYNTHETIC sample transaction files in sample_data/.

All rows are fictional. Dates fall after the demo workspace's 2026-08-31
reconciliation date so the import visibly updates cash and burn. The files
deliberately contain review cases: a duplicate row, an ambiguous date, a
transfer pair between entities, a refund, a financing inflow, a formula-like
text cell, a row with an invalid amount, and a payment matching an open
commitment.
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT = Path(__file__).resolve().parents[1] / "sample_data"
HEADER = ["Date", "Description", "Counterparty", "Amount", "Currency", "Program", "Milestone", "Category", "Account", "Reference"]

TEMPLATE_ROWS = [
    ["2026-01-15", "EXAMPLE - CRO invoice 001", "Example CRO", "-12,500.00", "USD", "PROGRAM-CODE", "Milestone name", "CRO services", "Your account name", "INV-001"],
    ["2026-01-20", "EXAMPLE - refund of reagent overcharge", "Example supplier", "300.00", "USD", "", "", "", "Your account name", "CN-17"],
]

LEDGER = [
    # date, description, counterparty, amount, currency, program, milestone, category, account, reference
    ["2026-09-05", "Payroll September", "Payroll provider (synthetic)", "-300,000.00", "HKD", "shared", "", "Personnel", "HK operating account (HKD)", "PAY-0926"],
    ["2026-09-10", "HX-101 GLP study milestone 1 (CRO)", "Demo CRO Ltd", "-96,000.00", "USD", "HX-101", "GLP target-animal safety study", "CRO services", "US operating account (USD)", "CRO-M1"],
    ["2026-09-12", "Bioanalytical assay transfer", "Demo Bioanalytics (synthetic)", "-150,000.00", "CNY", "HX-101", "GLP target-animal safety study", "Analytical work", "Shenzhen operating account (CNY)", "BA-0912"],
    ["2026-09-15", "Lab lease September", "Landlord (synthetic)", "-60,000.00", "CNY", "shared", "", "Facilities", "Shenzhen operating account (CNY)", "LEASE-09"],
    ["2026-09-18", "Insurance and software", "Various (synthetic)", "-8,000.00", "USD", "shared", "", "Overhead", "US operating account (USD)", "OH-09"],
    ["2026-09-22", "Screening reagents", "Reagent supplier (synthetic)", "-180,000.00", "HKD", "HX-205", "Lead compound selection", "Research", "HK operating account (HKD)", "RG-0922"],
    ["2026-09-25", "Refund of reagent overcharge", "Reagent supplier (synthetic)", "12,000.00", "HKD", "HX-205", "Lead compound selection", "Research", "HK operating account (HKD)", "CN-0925"],
    ["2026-09-28", "Intercompany transfer to US entity", "Own account", "-389,000.00", "HKD", "", "", "", "HK operating account (HKD)", "TRF-0928"],
    ["2026-09-29", "Intercompany transfer from HK entity", "Own account", "50,000.00", "USD", "", "", "", "US operating account (USD)", "TRF-0928"],
    ["2026-10-05", "Payroll October", "Payroll provider (synthetic)", "-300,000.00", "HKD", "shared", "", "Personnel", "HK operating account (HKD)", "PAY-1026"],
    ["2026-10-08", "Pilot tablet batch CMO deposit", "Demo CMO (synthetic)", "-40,000.00", "USD", "HX-101", "Pilot field efficacy study", "Manufacturing / CMC", "US operating account (USD)", "CDMO-1008"],
    ["06/10/2026", "Regulatory consultant (jurisdiction memo)", "Reg consultant (synthetic)", "-6,500.00", "USD", "HX-101", "", "Regulatory", "US operating account (USD)", "RC-1006"],
    ["2026-10-12", "Bank charges", "Bank (synthetic)", "-1,250.00", "HKD", "", "", "", "HK operating account (HKD)", "BC-1012"],
    ["2026-10-14", "Grant received (synthetic innovation grant)", "Government fund (synthetic)", "200,000.00", "HKD", "", "", "", "HK operating account (HKD)", "GR-1014"],
    ["2026-10-15", "Lab lease October", "Landlord (synthetic)", "-60,000.00", "CNY", "shared", "", "Facilities", "Shenzhen operating account (CNY)", "LEASE-10"],
    ["2026-10-16", "=HYPERLINK(\"https://example.org\",\"click\")", "Unknown (synthetic)", "-99.00", "USD", "", "", "", "US operating account (USD)", "X-1016"],
    ["2026-10-18", "Courier", "Courier (synthetic)", "abc", "HKD", "", "", "", "HK operating account (HKD)", "ERR-1018"],
    ["2026-10-20", "Screening reagents", "Reagent supplier (synthetic)", "-150,000.00", "HKD", "HX-205", "Lead compound selection", "Research", "HK operating account (HKD)", "RG-1020"],
    ["2026-09-05", "Payroll September", "Payroll provider (synthetic)", "-300,000.00", "HKD", "shared", "", "Personnel", "HK operating account (HKD)", "PAY-0926"],
]


def write_csv(path: Path, rows: list[list[str]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        w.writerows(rows)


def write_xlsx(path: Path, rows: list[list[str]], notes: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    ws.append([notes])
    ws["A1"].font = Font(bold=True)
    tx = wb.create_sheet("Transactions")
    tx.append(HEADER)
    for c in tx[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
    for r in rows:
        out = []
        for i, v in enumerate(r):
            if i == 3:
                try:
                    out.append(float(v.replace(",", "")))
                    continue
                except ValueError:
                    pass
            out.append(v)
        # store formula-looking text as text, not as a formula
        cell_row = tx.max_row + 1
        tx.append(out)
        for col, v in enumerate(out, start=1):
            if isinstance(v, str) and v.startswith("="):
                tx.cell(row=cell_row, column=col).data_type = "s"
    for col, width in zip("ABCDEFGHIJ", (12, 42, 30, 14, 9, 10, 30, 22, 34, 12)):
        tx.column_dimensions[col].width = width
    wb.save(path)


def main() -> None:
    OUT.mkdir(exist_ok=True)
    write_csv(OUT / "transactions_template.csv", TEMPLATE_ROWS)
    write_xlsx(OUT / "transactions_template.xlsx", TEMPLATE_ROWS, "Template. Replace the EXAMPLE rows. Negative amount = money out. Use HKD, CNY or USD.")
    hk = [r for r in LEDGER if "HK operating" in r[8]]
    write_csv(OUT / "synthetic_bank_export_hkd.csv", hk)
    write_xlsx(OUT / "synthetic_ledger_multi_entity.xlsx", LEDGER, "SYNTHETIC DEMO DATA - fictional transactions for the demo workspace (Sep-Oct 2026). Select the 'Transactions' sheet.")
    print("wrote", sorted(p.name for p in OUT.iterdir()))


if __name__ == "__main__":
    main()
